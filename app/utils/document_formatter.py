from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from fpdf import FPDF


BASE_DIR = Path(__file__).resolve().parents[2]
FONT_PATH = BASE_DIR / "assets" / "arial.ttf"


def sanitize_text(text: str) -> str:
    """Clean generated text before exporting."""
    if not text:
        return ""
    return text.replace("\r\n", "\n").strip()


def format_txt(text: str) -> bytes:
    """Create a TXT file."""
    clean_text = sanitize_text(text)
    return clean_text.encode("utf-8")


def format_docx(text: str, document_type: str) -> bytes:
    """Create a formatted DOCX document."""
    clean_text = sanitize_text(text)

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    title.alignment = 1

    run = title.add_run(document_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for paragraph in clean_text.split("\n"):
        if not paragraph.strip():
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)

        run = p.add_run(paragraph.strip())
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

    footer = section.footer.paragraphs[0]
    footer.alignment = 1

    footer_run = footer.add_run(
        "LegalEase | AI-generated draft | Review before use"
    )
    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(8)

    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return output.getvalue()


def format_pdf(text: str, document_type: str) -> bytes:
    """Create a Unicode-compatible PDF document."""
    clean_text = sanitize_text(text)

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)

    # Add Unicode Arial font
    pdf.add_font(
        "Arial",
        "",
        str(FONT_PATH)
    )

    pdf.add_font(
        "Arial",
        "B",
        str(FONT_PATH)
    )

    pdf.add_page()

    # Document title
    pdf.set_font("Arial", "B", 16)
    pdf.cell(
        0,
        10,
        document_type.upper(),
        align="C"
    )
    pdf.ln(14)

    # Document body
    pdf.set_font("Arial", "", 11)

    for paragraph in clean_text.split("\n"):
        if not paragraph.strip():
            pdf.ln(4)
            continue

        pdf.multi_cell(
            0,
            7,
            paragraph.strip()
        )
        pdf.ln(2)

    # Footer
    pdf.set_y(-15)
    pdf.set_font("Arial", "", 8)

    pdf.cell(
        0,
        10,
        "LegalEase | AI-generated draft | Review before use",
        align="C"
    )

    return bytes(pdf.output())